"""Excel and PDF export of the stored 12-month plan (US-18, BE-21, BE-22).

File names contain the date and the plan version, e.g.
budce_plani_v3_2026-10-03.xlsx
"""

from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Same columns as the dashboard / specification (§28).
TABLE_COLUMNS = [
    ("month_name", "Ay"),
    ("income", "Gəlir"),
    ("housing", "Kirayə"),
    ("restaurant", "Restoran"),
    ("entertainment", "Əyləncə"),
    ("food", "Qida"),
    ("utilities", "Kommunal"),
    ("transport", "Nəqliyyat"),
    ("credit", "Kredit"),
    ("other", "Digər"),
    ("savings", "Yığım"),
    ("balance", "Qalıq"),
]
COMPARISON_COLUMNS = [
    ("category_name", "Kateqoriya"),
    ("percentage", "%"),
    ("current_monthly_amount", "Hazırkı aylıq"),
    ("recommended_monthly_amount", "Tövsiyə olunan aylıq"),
    ("annual_amount", "İllik"),
    ("status", "Status"),
    ("ai_recommendation", "AI tövsiyəsi"),
]

FONT_DIR = Path(__file__).resolve().parent / "fonts"
_FONTS_READY = False


def _filename(session, ext):
    date = timezone.localdate().isoformat()
    return f"budce_plani_v{session.plan_version}_{date}.{ext}"


def _num(value):
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _annual_row(session):
    totals = session.annual_totals or {}
    return ["İllik cəmi"] + [_num(totals.get(f"total_{key}")) for key, _ in TABLE_COLUMNS[1:]]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def build_excel_response(session):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "12 aylıq plan"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="FFE8D6")

    ws.append(["SmartBudget AI — 12 aylıq büdcə planı"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"İstifadəçi: {session.user.get_full_name() or session.user.email}"])
    ws.append([f"Maliyyə vəziyyəti: {session.financial_status}"])
    ws.append([f"Tövsiyə olunan aylıq yığım: {_num(session.recommended_monthly_savings)} AZN"])
    ws.append([f"Tövsiyə olunan illik yığım: {_num(session.recommended_annual_savings)} AZN"])
    ws.append([f"Plan versiyası: {session.plan_version} · Tarix: {timezone.localdate().isoformat()}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append([label + ("" if key == "month_name" else " (AZN)") for key, label in TABLE_COLUMNS] + ["Qeyd"])
    for row in session.monthly_table or []:
        ws.append([row.get("month_name")] + [_num(row.get(key)) for key, _ in TABLE_COLUMNS[1:]] + [row.get("note", "")])
    ws.append(_annual_row(session))
    for cell in ws[header_row]:
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for cell in ws[ws.max_row]:
        cell.font = bold
    for col in range(1, len(TABLE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(len(TABLE_COLUMNS) + 1)].width = 90
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=2, max_col=len(TABLE_COLUMNS)):
        for cell in row:
            cell.number_format = "#,##0.00"

    ws2 = wb.create_sheet("Büdcə müqayisəsi")
    ws2.append([label for _, label in COMPARISON_COLUMNS])
    for cell in ws2[1]:
        cell.font = bold
        cell.fill = head_fill
    for row in session.budget_comparison or []:
        ws2.append([
            row.get(key) if key in ("category_name", "status", "ai_recommendation") else _num(row.get(key))
            for key, _ in COMPARISON_COLUMNS
        ])
    for col, width in zip("ABCDEFG", (24, 8, 16, 20, 14, 18, 70)):
        ws2.column_dimensions[col].width = width

    goals = session.savings_goals_breakdown or []
    if goals:
        ws3 = wb.create_sheet("Yığım məqsədləri")
        ws3.append(["Məqsəd", "Prioritet", "Hədəf (AZN)", "Artıq yığılıb (AZN)", "Gedişat %", "Müddət (ay)",
                    "Lazımi aylıq orta (AZN)", "12 ayda planlanan (AZN)", "12 aydan sonra (AZN)",
                    "Müddətin sonunda gözlənilən (AZN)", "Vaxtında?", "Qeyd"])
        for cell in ws3[1]:
            cell.font = bold
            cell.fill = head_fill
        for g in goals:
            ws3.append([g.get("goal_name"), g.get("priority"), _num(g.get("target_amount")), _num(g.get("saved_amount")),
                        _num(g.get("progress_percentage")), g.get("deadline_months"),
                        _num(g.get("required_monthly_average")), _num(g.get("planned_contribution_12m")),
                        _num(g.get("projected_amount_12m")), _num(g.get("expected_amount_at_deadline")),
                        "Bəli" if g.get("on_track") else "Xeyr", g.get("note", "")])
        for col, width in zip("ABCDEFGHIJKL", (34, 16, 13, 18, 11, 12, 20, 20, 18, 26, 10, 70)):
            ws3.column_dimensions[col].width = width

        # Monthly Savings split by goal: Yığım = sum of the goal contributions.
        ws4 = wb.create_sheet("Məqsədlər üzrə yığım")
        ws4.append(["Ay"] + [g.get("goal_name") for g in goals] + ["Yığım cəmi (AZN)", "Qalıq (AZN)"])
        for cell in ws4[1]:
            cell.font = bold
            cell.fill = head_fill
        for i, row in enumerate(session.monthly_table or []):
            parts = [_num(c) for c in (row.get("goal_contributions") or [])]
            ws4.append([row.get("month_name")] + parts + [_num(row.get("savings")), _num(row.get("balance"))])
        ws4.append(["İllik cəmi"] + [_num(g.get("planned_contribution_12m")) for g in goals]
                   + [_num((session.annual_totals or {}).get("total_savings")),
                      _num((session.annual_totals or {}).get("total_balance"))])
        for cell in ws4[ws4.max_row]:
            cell.font = bold
        for col in range(1, len(goals) + 4):
            ws4.column_dimensions[get_column_letter(col)].width = 22

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{_filename(session, "xlsx")}"'
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    wb.save(response)
    return response


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _register_fonts():
    """DejaVu supports Azerbaijani letters (ə, ğ, ı, ş, ç, ö, ü); Helvetica does not."""
    global _FONTS_READY
    if not _FONTS_READY:
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSansCondensed.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSansCondensed-Bold.ttf")))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                      italic="DejaVu", boldItalic="DejaVu-Bold")
        _FONTS_READY = True


def build_pdf_response(session):
    _register_fonts()
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(session, "pdf")}"'
    response["Access-Control-Expose-Headers"] = "Content-Disposition"

    doc = SimpleDocTemplate(response, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="SmartBudget AI — büdcə planı")
    title = ParagraphStyle("t", fontName="DejaVu-Bold", fontSize=15, leading=19, spaceAfter=4)
    body = ParagraphStyle("b", fontName="DejaVu", fontSize=9, leading=12)
    small = ParagraphStyle("s", fontName="DejaVu", fontSize=7.5, leading=9.5)
    h2 = ParagraphStyle("h", fontName="DejaVu-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)

    story = [
        Paragraph("SmartBudget AI — 12 aylıq büdcə planı", title),
        Paragraph(
            f"İstifadəçi: {escape(session.user.get_full_name() or session.user.email)} · "
            f"Plan versiyası: {session.plan_version} · Tarix: {timezone.localdate().isoformat()}", body),
        Paragraph(
            f"Maliyyə vəziyyəti: <b>{escape(session.financial_status)}</b> — {escape(session.financial_status_description)}", body),
        Paragraph(
            f"Tövsiyə olunan aylıq yığım: <b>{_num(session.recommended_monthly_savings):,.2f} AZN</b> · "
            f"illik: <b>{_num(session.recommended_annual_savings):,.2f} AZN</b>", body),
        Paragraph(escape(session.ai_response_text or ""), body),
        *([Paragraph("<b>Diqqət:</b> büdcə mövcud gəlir və məcburi öhdəliklərlə mümkün deyil — "
                     "qırmızı qalıq aylıq kəsiri göstərir.", ParagraphStyle("w", parent=body, textColor=colors.HexColor("#DC2626")))]
          if any(_num(r.get("balance")) < 0 for r in (session.monthly_table or [])) else []),
        Spacer(1, 4 * mm),
        Paragraph("12 aylıq plan (AZN)", h2),
    ]

    data = [[label for _, label in TABLE_COLUMNS]]
    negative_rows = []
    for i, row in enumerate(session.monthly_table or [], start=1):
        data.append([row.get("month_name")] + [f"{_num(row.get(k)):,.2f}" for k, _ in TABLE_COLUMNS[1:]])
        if _num(row.get("balance")) < 0:
            negative_rows.append(i)
    data.append([_annual_row(session)[0]] + [f"{v:,.2f}" for v in _annual_row(session)[1:]])
    table = Table(data, repeatRows=1)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
        ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "DejaVu-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFE8D6")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F4F6")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
    ]
    for r in negative_rows:
        style.append(("TEXTCOLOR", (-1, r), (-1, r), colors.HexColor("#DC2626")))
    table.setStyle(TableStyle(style))
    story.append(table)

    comparison = session.budget_comparison or []
    if comparison:
        story.append(Paragraph("Hazırkı və tövsiyə olunan büdcə", h2))
        rows = [[label for _, label in COMPARISON_COLUMNS]]
        for row in comparison:
            rows.append([
                Paragraph(escape(str(row.get("category_name", ""))), small),
                f"{_num(row.get('percentage')):.2f}",
                f"{_num(row.get('current_monthly_amount')):,.2f}",
                f"{_num(row.get('recommended_monthly_amount')):,.2f}",
                f"{_num(row.get('annual_amount')):,.2f}",
                Paragraph(escape(str(row.get("status", ""))), small),
                Paragraph(escape(str(row.get("ai_recommendation", ""))), small),
            ])
        ct = Table(rows, colWidths=[40 * mm, 14 * mm, 26 * mm, 32 * mm, 26 * mm, 30 * mm, 100 * mm], repeatRows=1)
        ct.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
            ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFE8D6")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(ct)

    goals = session.savings_goals_breakdown or []
    if goals:
        story.append(Paragraph("Yığım məqsədləri", h2))
        for g in goals:
            story.append(Paragraph(
                f"<b>{escape(str(g.get('goal_name')))}</b> ({escape(str(g.get('priority')))}): hədəf "
                f"{_num(g.get('target_amount')):,.2f} AZN, artıq yığılıb {_num(g.get('saved_amount')):,.2f} AZN "
                f"({_num(g.get('progress_percentage')):.1f}%), müddət {g.get('deadline_months')} ay, 12 ayda planlanan "
                f"{_num(g.get('planned_contribution_12m')):,.2f} AZN (orta aylıq {_num(g.get('recommended_monthly_saving')):,.2f} AZN). "
                f"{escape(str(g.get('note', '')))}", body))
        split = [["Ay"] + [Paragraph(escape(str(g.get("goal_name"))), small) for g in goals] + ["Yığım", "Qalıq"]]
        for row in session.monthly_table or []:
            split.append([row.get("month_name")] + [f"{_num(c):,.2f}" for c in (row.get("goal_contributions") or [])]
                         + [f"{_num(row.get('savings')):,.2f}", f"{_num(row.get('balance')):,.2f}"])
        gt = Table(split, repeatRows=1)
        gt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
            ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFE8D6")),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
        ]))
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph("Yığımın məqsədlər üzrə aylıq bölgüsü (AZN)", body))
        story.append(gt)

    story.append(Paragraph("Aylar üzrə izah", h2))
    for row in session.monthly_table or []:
        story.append(Paragraph(f"<b>{escape(str(row.get('month_name')))}:</b> {escape(str(row.get('note', '')))}", small))

    doc.build(story)
    return response
